#!/usr/bin/env python3
"""
SIGAI-SES - Generador de Documentos Word Profesionales (APA)
============================================================
Convierte archivos Markdown a Word con formato APA profesional:
  - Portada corporativa SIGAI-SES
  - Tabla de contenido automatica
  - Tipografia Times New Roman 12pt
  - Interlineado 2.0 (APA)
  - Margenes 2.54 cm
  - Numeracion de paginas
  - Encabezados con titulo
  - Tablas con formato profesional
  - Bloques de codigo con fondo gris

Uso: python generar_word_apa.py
"""

import os
import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import markdown
from html.parser import HTMLParser

# Configuracion
SCRIPT_DIR = Path(__file__).parent
OUTPUT_DIR = SCRIPT_DIR / "Word_Entregables"
IMAGES_DIR = SCRIPT_DIR / "images"

# Colores corporativos SES
SES_BLUE = RGBColor(0x1a, 0x3c, 0x6e)
SES_ORANGE = RGBColor(0xFF, 0x6B, 0x35)
SES_GRAY = RGBColor(0x66, 0x66, 0x66)
SES_LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)


class MarkdownToDocx:
    """Convierte Markdown a Word con formato APA profesional."""

    def __init__(self, title: str, author: str = "Wilson Ortiz"):
        self.title = title
        self.author = author
        self.doc = Document()
        self._setup_styles()

    def _setup_styles(self):
        """Configura estilos del documento APA."""
        style = self.doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)
        font.color.rgb = RGBColor(0, 0, 0)

        pf = style.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.first_line_indent = Cm(1.27)

        # Margenes APA: 2.54 cm por lado
        for section in self.doc.sections:
            section.top_margin = Cm(2.54)
            section.bottom_margin = Cm(2.54)
            section.left_margin = Cm(2.54)
            section.right_margin = Cm(2.54)

        # Configurar encabezado
        header = self.doc.sections[0].header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = hp.add_run(f"SIGAI-SES — SES — Seguridad Electrónica | {self.title}")
        run.font.size = Pt(9)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = SES_GRAY

        # Configurar pie de pagina (numeracion)
        footer = self.doc.sections[0].footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run()
        fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
        run._r.append(fldChar1)
        run2 = fp.add_run()
        instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
        run2._r.append(instrText)
        run3 = fp.add_run()
        fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
        run3._r.append(fldChar2)
        for r in [run, run2, run3]:
            r.font.size = Pt(9)
            r.font.name = 'Times New Roman'

    def add_cover_page(self):
        """Agrega portada profesional APA."""
        # Espaciado superior
        for _ in range(4):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)

        # Nombre del proyecto
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run("SISTEMA INTEGRAL DE GESTIÓN DE ACTIVOS E INVENTARIO")
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.name = 'Times New Roman'
        run.font.color.rgb = SES_BLUE

        # SIGAI-SES
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(30)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run("SIGAI-SES")
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.name = 'Times New Roman'
        run.font.color.rgb = SES_BLUE

        # Linea decorativa
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(20)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run("━" * 50)
        run.font.size = Pt(10)
        run.font.color.rgb = SES_ORANGE

        # Titulo del documento
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(40)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(self.title.upper().replace("_", " ").replace("-", " "))
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.name = 'Times New Roman'
        run.font.color.rgb = SES_BLUE

        # Autor
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(self.author)
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'

        # Cargo
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run("Pasante SENA")
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'

        # Programa
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(20)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run("Tecnología en Análisis y Desarrollo de Software")
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'

        # Empresa
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run("SES — Seguridad Electrónica")
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.name = 'Times New Roman'
        run.font.color.rgb = SES_BLUE

        # Fecha
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(40)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run("Julio 2026")
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = SES_GRAY

        # Cuadro confidencial
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run("Documento Confidencial")
        run.font.size = Pt(9)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = SES_GRAY

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run("SES — Seguridad Electrónica — Todos los derechos reservados")
        run.font.size = Pt(9)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = SES_GRAY

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run("Versión 1.0.0 · Julio 2026")
        run.font.size = Pt(9)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = SES_GRAY

        # Salto de pagina
        self.doc.add_page_break()

    def add_toc_placeholder(self):
        """Agrega marcador para tabla de contenido (se actualiza en Word)."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(20)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run("TABLA DE CONTENIDO")
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.name = 'Times New Roman'
        run.font.color.rgb = SES_BLUE

        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("[Actualizar tabla de contenido en Word: clic derecho → Actualizar campo]")
        run.font.size = Pt(10)
        run.font.italic = True
        run.font.name = 'Times New Roman'
        run.font.color.rgb = SES_GRAY

        # Insertar campo TOC de Word
        p2 = self.doc.add_paragraph()
        p2.paragraph_format.first_line_indent = Cm(0)
        run = p2.add_run()
        fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
        run._r.append(fldChar1)
        run2 = p2.add_run()
        instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>')
        run2._r.append(instrText)
        run3 = p2.add_run()
        fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
        run3._r.append(fldChar2)
        run4 = p2.add_run()
        fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
        run4._r.append(fldChar3)

        self.doc.add_page_break()

    def _add_heading(self, text: str, level: int):
        """Agrega encabezado con formato APA."""
        h = self.doc.add_heading(text, level=level)
        h.paragraph_format.first_line_indent = Cm(0)
        h.paragraph_format.space_before = Pt(18 if level == 1 else 12 if level == 2 else 8)
        h.paragraph_format.space_after = Pt(10 if level == 1 else 8)
        for run in h.runs:
            run.font.name = 'Times New Roman'
            run.font.color.rgb = SES_BLUE
            if level == 1:
                run.font.size = Pt(14)
            elif level == 2:
                run.font.size = Pt(13)
            else:
                run.font.size = Pt(12)
        return h

    def _add_paragraph(self, text: str, bold=False, italic=False, indent=True):
        """Agrega parrafo con formato APA."""
        p = self.doc.add_paragraph()
        if not indent:
            p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = bold
        run.font.italic = italic
        return p

    def _add_code_block(self, code: str):
        """Agrega bloque de codigo con fondo gris."""
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)

        # Fondo gris via sombra
        pPr = p._p.get_or_add_pPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F0F0F0"/>')
        pPr.append(shd)

        run = p.add_run(code)
        run.font.name = 'Courier New'
        run.font.size = Pt(10)
        return p

    def _add_table(self, headers: list, rows: list):
        """Agrega tabla con formato profesional."""
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        # Encabezados
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(header.strip())
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Fondo azul oscuro
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="1A3C6E"/>')
            tcPr.append(shd)

        # Filas de datos
        num_cols = len(headers)
        for r_idx, row in enumerate(rows):
            for c_idx, cell_text in enumerate(row):
                if c_idx >= num_cols:
                    break
                cell = table.rows[r_idx + 1].cells[c_idx]
                cell.text = ""
                p = cell.paragraphs[0]
                # Limpiar HTML y Markdown de las celdas (badges, etc.)
                clean_text = str(cell_text).strip()
                # Extraer nombre de tecnologia de badges shields.io
                badge_match = re.search(r'badge/-([A-Za-z0-9+._% -]+?)-[0-9A-Fa-f]{3,6}(?:\?|$)', clean_text)
                if not badge_match:
                    badge_match = re.search(r'badge/([A-Za-z0-9+._% -]+?)-[0-9A-Fa-f]{3,6}(?:\?|$)', clean_text)
                if badge_match:
                    name = badge_match.group(1)
                    name = name.replace('-', ' ').replace('+', '+').replace('_', ' ')
                    name = name.replace('%20', ' ').replace('%2F', '/')
                    # Limpiar dobles espacios
                    name = re.sub(r'\s+', ' ', name).strip()
                    clean_text = name
                else:
                    clean_text = re.sub(r'<img[^>]*>', '', clean_text)
                    clean_text = re.sub(r'<br\s*/?>', '\n', clean_text)
                    clean_text = re.sub(r'<[^>]+>', '', clean_text)
                    # Limpiar markdown images ![alt](url)
                    clean_text = re.sub(r'!\[([^\]]*)\]\([^)]+\)', r'\1', clean_text)
                run = p.add_run(clean_text)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                # Fondo alternado
                if r_idx % 2 == 1:
                    tc = cell._tc
                    tcPr = tc.get_or_add_tcPr()
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F0F4F8"/>')
                    tcPr.append(shd)

        # Espacio despues de tabla
        self.doc.add_paragraph()
        return table

    def _parse_markdown_line(self, line: str):
        """Parsea una linea de markdown y la agrega al documento."""
        line = line.rstrip()

        # Linea vacia
        if not line:
            return None

        # Encabezados
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            text = line.lstrip('#').strip()
            if level <= 3 and text:
                self._add_heading(text, min(level, 3))
                return 'heading'

        # Bloques de codigo
        if line.startswith('```'):
            return 'code_boundary'

        # Tablas
        if line.startswith('|') and '|' in line[1:]:
            return 'table_row'

        # Separadores
        if line.startswith('---') or line.startswith('==='):
            return 'separator'

        # Listas
        if re.match(r'^[-*+]\s', line):
            text = re.sub(r'^[-*+]\s', '', line).strip()
            p = self.doc.add_paragraph(text, style='List Bullet')
            p.paragraph_format.first_line_indent = Cm(0)
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)
            return 'list'

        if re.match(r'^\d+\.\s', line):
            text = re.sub(r'^\d+\.\s', '', line).strip()
            p = self.doc.add_paragraph(text, style='List Number')
            p.paragraph_format.first_line_indent = Cm(0)
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)
            return 'list'

        # Blockquotes
        if line.startswith('>'):
            text = line.lstrip('>').strip()
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.27)
            p.paragraph_format.first_line_indent = Cm(0)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.italic = True
            run.font.color.rgb = SES_GRAY
            return 'quote'

        # Imagenes
        img_match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line)
        if img_match:
            alt, src = img_match.groups()
            # Buscar imagen en multiples ubicaciones relativas al .md
            img_found = False
            search_paths = [
                self.md_dir / src,                          # relativo al .md
                self.md_dir / "images" / Path(src).name,    # images/ subfolder
                SCRIPT_DIR / src,                           # relativo a ENTREGABLES
                SCRIPT_DIR / "images" / Path(src).name,     # images/ global
            ]
            # Buscar tambien por nombre de archivo en todas las carpetas images/
            img_name = Path(src).name
            for img_dir in self.md_dir.rglob("images"):
                search_paths.append(img_dir / img_name)

            for img_path in search_paths:
                if img_path.exists() and img_path.suffix.lower() in ['.png', '.jpg', '.jpeg', '.gif', '.bmp']:
                    try:
                        p = self.doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.first_line_indent = Cm(0)
                        p.paragraph_format.space_before = Pt(6)
                        p.paragraph_format.space_after = Pt(6)
                        run = p.add_run()
                        run.add_picture(str(img_path), width=Inches(5.5))
                        img_found = True
                        break
                    except Exception as e:
                        pass

            if not img_found:
                # Si no encontro la imagen, agregar referencia como texto
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
                run = p.add_run(f"[Imagen: {alt or src}]")
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                run.font.italic = True
                run.font.color.rgb = SES_GRAY

            return 'image'

        # HTML basico (badges, etc.) - extraer texto
        if '<' in line:
            # Detectar etiquetas <img> y embedirlas
            img_html = re.search(r'<img[^>]+src="([^"]+)"[^>]*>', line)
            if img_html:
                src = img_html.group(1)
                # Solo procesar imagenes locales, ignorar URLs externas (badges)
                if not src.startswith('http'):
                    # Buscar imagen local
                    img_found = False
                    search_paths = [
                        self.md_dir / src,
                        self.md_dir / "images" / Path(src).name,
                        SCRIPT_DIR / src,
                        SCRIPT_DIR / "images" / Path(src).name,
                    ]
                    img_name = Path(src).name
                    for img_dir in self.md_dir.rglob("images"):
                        search_paths.append(img_dir / img_name)

                    for img_path in search_paths:
                        if img_path.exists() and img_path.suffix.lower() in ['.png', '.jpg', '.jpeg', '.gif', '.bmp']:
                            try:
                                p = self.doc.add_paragraph()
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                p.paragraph_format.first_line_indent = Cm(0)
                                p.paragraph_format.space_before = Pt(6)
                                p.paragraph_format.space_after = Pt(6)
                                run = p.add_run()
                                run.add_picture(str(img_path), width=Inches(5.5))
                                img_found = True
                                break
                            except:
                                pass
                    if img_found:
                        return 'image_html'
                else:
                    # URL externa (badges, etc.) - ignorar completamente
                    return 'html_skip'

            # Remover tags HTML
            clean = re.sub(r'<[^>]+>', '', line).strip()
            if clean and clean != line.strip():
                self._add_paragraph(clean)
                return 'html'
            # Si es solo HTML sin texto util, ignorar
            return 'html_skip'

        # Parrafo normal con formato inline
        text = line.strip()
        if text:
            p = self.doc.add_paragraph()
            # Procesar formato inline
            self._add_inline_formatting(p, text)
            return 'paragraph'

        return None

    def _add_inline_formatting(self, paragraph, text: str):
        """Agrega texto con formato inline (bold, italic, code, links)."""
        # Simplificar: remover markdown badges y shields.io
        text = re.sub(r'!\[[^\]]*\]\([^)]+\)', '', text)  # remover imagenes
        text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)  # links a texto
        text = re.sub(r'`([^`]+)`', r'\1', text)  # inline code a texto

        # Remover badges HTML restantes
        text = re.sub(r'<img[^>]*>', '', text)
        text = re.sub(r'<br\s*/?>', '\n', text)
        text = re.sub(r'<[^>]+>', '', text)

        text = text.strip()
        if not text:
            return

        # Detectar bold y italic
        parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|__[^_]+__|_[^_]+_)', text)
        for part in parts:
            if not part:
                continue
            run = paragraph.add_run()
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)

            if part.startswith('**') and part.endswith('**'):
                run.text = part[2:-2]
                run.font.bold = True
            elif part.startswith('__') and part.endswith('__'):
                run.text = part[2:-2]
                run.font.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run.text = part[1:-1]
                run.font.italic = True
            elif part.startswith('_') and part.endswith('_'):
                run.text = part[1:-1]
                run.font.italic = True
            else:
                run.text = part

    def convert_md_to_docx(self, md_path: Path, output_path: Path):
        """Convierte un archivo Markdown a Word profesional."""
        self.doc = Document()
        self._setup_styles()

        # Directorio base del archivo .md (para buscar imagenes relativas)
        self.md_dir = md_path.parent

        with open(md_path, 'r', encoding='utf-8') as f:
            content = f.read()

        # Portada
        self.add_cover_page()

        # Tabla de contenido
        self.add_toc_placeholder()

        # Procesar contenido
        lines = content.split('\n')
        in_code_block = False
        code_buffer = []
        in_table = False
        table_headers = []
        table_rows = []
        frontmatter_done = False

        i = 0
        while i < len(lines):
            line = lines[i]

            # Saltar frontmatter YAML
            if i == 0 and line.strip() == '---':
                i += 1
                while i < len(lines) and lines[i].strip() != '---':
                    i += 1
                i += 1
                frontmatter_done = True
                continue

            # Bloques de codigo
            if line.strip().startswith('```'):
                if in_code_block:
                    # Fin del bloque
                    self._add_code_block('\n'.join(code_buffer))
                    code_buffer = []
                    in_code_block = False
                else:
                    # Inicio del bloque
                    in_code_block = True
                i += 1
                continue

            if in_code_block:
                code_buffer.append(line)
                i += 1
                continue

            # Tablas
            if line.strip().startswith('|') and '|' in line.strip()[1:]:
                cells = [c.strip() for c in line.strip().split('|')[1:-1]]

                # Detectar separador de tabla
                if all(re.match(r'^[-:]+$', c) for c in cells):
                    i += 1
                    continue

                if not in_table:
                    in_table = True
                    table_headers = cells
                else:
                    table_rows.append(cells)

                # Verificar si la siguiente linea es tabla
                next_is_table = False
                if i + 1 < len(lines):
                    next_line = lines[i + 1].strip()
                    if next_line.startswith('|') and '|' in next_line[1:]:
                        next_is_table = True

                if not next_is_table and in_table:
                    self._add_table(table_headers, table_rows)
                    in_table = False
                    table_headers = []
                    table_rows = []

                i += 1
                continue

            # Procesar linea normal
            result = self._parse_markdown_line(line)
            i += 1

        # Guardar
        output_path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(output_path))


def main():
    """Funcion principal."""
    script_dir = Path(__file__).parent
    output_dir = script_dir / "Word_Entregables"
    output_dir.mkdir(exist_ok=True)

    print("=" * 50)
    print("  SIGAI-SES - Generador Word APA")
    print("  SES — Seguridad Electrónica")
    print("=" * 50)
    print()

    # Buscar archivos .md
    md_files = []
    for md in script_dir.rglob("*.md"):
        # Excluir carpetas de salida y otras
        if any(ex in str(md) for ex in ['Word_Entregables', 'DOCS_EDITABLES', 'PDF-profesionales', 'images', 'tmp_md', 'styles']):
            continue
        md_files.append(md)

    md_files.sort()

    total_ok = 0
    total_error = 0
    errors = []

    for md_path in md_files:
        name = md_path.stem
        rel_path = md_path.relative_to(script_dir)
        output_path = output_dir / f"{name}.docx"

        # Saltar si ya existe y es reciente
        if output_path.exists() and output_path.stat().st_mtime > md_path.stat().st_mtime:
            print(f"  SKIP: {name} (ya existe)")
            total_ok += 1
            continue

        print(f"  Procesando: {rel_path}")

        try:
            converter = MarkdownToDocx(title=name)
            converter.convert_md_to_docx(md_path, output_path)
            size_kb = output_path.stat().st_size // 1024
            print(f"    OK: {name}.docx ({size_kb} KB)")
            total_ok += 1
        except Exception as e:
            print(f"    ERROR: {name} - {e}")
            total_error += 1
            errors.append(str(rel_path))

    print()
    print("=" * 50)
    print("  RESULTADO")
    print("=" * 50)
    print(f"  Word generados: {total_ok}")
    if total_error > 0:
        print(f"  Errores: {total_error}")
        for e in errors:
            print(f"    - {e}")
    print(f"  Ubicacion: {output_dir}")
    print()
    print("  Para actualizar tabla de contenido en Word:")
    print("    1. Abrir el documento")
    print("    2. Ctrl+A (seleccionar todo)")
    print("    3. F9 (actualizar campos)")
    print()


if __name__ == "__main__":
    main()
